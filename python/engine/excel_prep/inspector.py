"""只读体检：扫一张表，把「跑之前就该知道的问题」列出来（不动任何文件）。

**只读是这一层的全部价值。** 上传后先给结论、再让用户决定要不要整理，
用户才敢把体检挂在「上传即自动跑」的流程上。测试里单独钉了一条
「输入文件一个字节都没变」。

**行号一律是原文件的行号。** 用户是拿着原文件去核对的；整理之后行号会漂，
所以报告与台账都锚在原文件上 —— 用户照着「第 7 行」能在自己的表里找到那一行。

判定的口径（「为空」「相等」）全部来自 `values`，与整理层共用一套。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook

from .. import template
from . import rules, values


@dataclass(frozen=True)
class Issue:
    """一类问题占一条。

    `rows` 是它在原文件里涉及的行号；`cols` 是涉及的列序号（1 起）。
    表头类问题的主语是**列**（「第 3 列没列名」），行级问题的主语是行 ——
    整理要照着这两个坐标去删行/删列，报告也要照着它把用户的视线引到格子上。
    """

    kind: str
    severity: str
    rows: tuple[int, ...]
    detail: str
    cols: tuple[int, ...] = ()


@dataclass
class Report:
    path: str
    sheet: str
    #: 原文件里的表头行号。表头不在第 1 行时，`detect_columns`（读第 1 行）会读到空列名
    header_row: int
    columns: list[Any]
    key_col: str
    #: 表头之后的行数，与 `excel.count_rows` 同一口径（含空行）——
    #: 它要能解释「上传后说 9 行，跑起来只有 6 行有内容」
    data_rows: int
    issues: list[Issue] = field(default_factory=list)


def inspect(path: str, key_col: str = "") -> Report:
    """体检一张表。

    `key_col` 为空表示用户还没选主键列 —— 主键相关的判定整体跳过，不抛异常
    （刚上传就问结论的场景不能炸）。列名写错时抛 `ValueError`：静默当成
    「没有主键」，用户会拿着一份少了重复判定的体检结论去整理。
    """
    sheet, grid = read_grid(path)
    header_row = _first_content_row(grid)
    columns = list(grid[header_row - 1]) if grid else []
    key_idx = _key_index(columns, key_col)
    key_name = values.cell_text(key_col)

    issues: list[Issue] = []
    issues += _header_issues(grid, header_row, columns)
    issues += _leading_blank_issues(header_row)

    data = [(no, grid[no - 1]) for no in range(header_row + 1, len(grid) + 1)]
    dropped = _row_level_issues(data, key_idx, key_name, issues)
    _duplicate_issues(data, key_idx, dropped, issues)
    _unparsable_issues(grid, header_row, columns, issues)

    return Report(
        path=str(path),
        sheet=sheet,
        header_row=header_row,
        columns=columns,
        key_col=key_col,
        data_rows=max(len(grid) - header_row, 0),
        issues=issues,
    )


# ── 读 ────────────────────────────────────────────────────────────────


def read_grid(path: str) -> tuple[str, list[list[Any]]]:
    """整张表读成 `grid[i] = Excel 第 i+1 行`，每行补齐到等宽。

    `min_row=1` 是**必须显式写**的：表头落在第 3 行时，openpyxl 会把
    `dimension` 裁成 `A3:C4`，只读模式的 `iter_rows` 仍从第 1 行补齐返回 ——
    显式写出来才不依赖这个隐式行为，行号也才对得上原文件。
    """
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        max_row = ws.max_row or 0
        raw = [list(r) for r in ws.iter_rows(min_row=1, max_row=max_row, values_only=True)] if max_row else []
        width = max([ws.max_column or 0] + [len(r) for r in raw] + [0])
        grid = [list(r) + [None] * (width - len(r)) for r in raw]
        return ws.title, grid
    finally:
        wb.close()


def _first_content_row(grid: list[list[Any]]) -> int:
    """表头 = 第一个有内容的行。全空表退回 1，免得后面按行号切的 range 全乱。"""
    for no, row in enumerate(grid, start=1):
        if any(not values.is_blank(v) for v in row):
            return no
    return 1


def _key_index(columns: list[Any], key_col: str) -> int | None:
    """主键列的下标；`None` = 用户还没选主键列，主键相关的判定整体跳过。"""
    if values.is_blank(key_col):
        return None
    wanted = values.cell_text(key_col)
    for idx, name in enumerate(columns):
        if values.cell_text(name) == wanted:
            return idx
    raise ValueError(f"未找到列「{key_col}」，可用列：{columns}")


# ── 表头 ──────────────────────────────────────────────────────────────


def _header_issues(grid: list[list[Any]], header_row: int, columns: list[Any]) -> list[Issue]:
    """表头自身的三类毛病，按执行顺序返回。

    列名没定下来，后面按列名跑的一切都没意义 —— 所以它们排在最前。
    """
    if not columns:
        return []

    out: list[Issue] = []

    padded = [(idx + 1, c) for idx, c in enumerate(columns) if isinstance(c, str) and c != values.normalize_text(c)]
    if padded:
        out.append(
            Issue(
                kind="header_whitespace",
                severity=rules.severity("header_whitespace"),
                rows=(header_row,),
                cols=tuple(no for no, _ in padded),
                detail=rules.header_whitespace_detail([values.cell_text(c) for _, c in padded]),
            )
        )

    for idx, name in enumerate(columns):
        if not values.is_blank(name):
            continue
        col_no = idx + 1
        # 「整列为空」要连着数据区一起看：表头单元格自己就是空的，只看它会恒判为空
        empty = all(values.is_blank(row[idx]) for row in grid[header_row:])
        out.append(
            Issue(
                kind="blank_header",
                severity=rules.blank_header_severity(empty),
                rows=(header_row,),
                cols=(col_no,),
                detail=rules.blank_header_detail(col_no, empty),
            )
        )

    seen: dict[str, int] = {}
    for idx, name in enumerate(columns):
        label = values.cell_text(name)
        if not label:  # 没列名的列不算重名 —— 那是 blank_header 的事
            continue
        first = seen.setdefault(label, idx + 1)
        if first != idx + 1:
            out.append(
                Issue(
                    kind="duplicate_header",
                    severity=rules.severity("duplicate_header"),
                    rows=(header_row,),
                    cols=(idx + 1,),
                    detail=rules.duplicate_header_detail(idx + 1, first, label),
                )
            )

    return out


def _leading_blank_issues(header_row: int) -> list[Issue]:
    if header_row <= 1:
        return []
    rows = tuple(range(1, header_row))
    return [
        Issue(
            kind="leading_blank_rows",
            severity=rules.severity("leading_blank_rows"),
            rows=rows,
            detail=rules.leading_blank_rows_detail(len(rows)),
        )
    ]


# ── 行级 ──────────────────────────────────────────────────────────────


def _row_level_issues(data: list[tuple[int, list[Any]]], key_idx: int | None, key_name: str, issues: list[Issue]) -> set[int]:
    """按执行顺序判「整行为空」「主键有空白」「主键为空」，返回被丢掉的行的行号。

    丢掉的集合要往外传：`duplicate_row` / `duplicate_key` 只看活下来的行。
    否则第 8 行（主键为空）会再以「主键重复」的名义报一次，
    用户回头在表里找不到它对应的是哪一条。
    """
    dropped: set[int] = set()

    blank_rows = [no for no, row in data if all(values.is_blank(v) for v in row)]
    if blank_rows:
        dropped.update(blank_rows)
        issues.append(
            Issue(
                kind="blank_row",
                severity=rules.severity("blank_row"),
                rows=tuple(blank_rows),
                detail=rules.blank_row_detail(),
            )
        )

    if key_idx is None:
        return dropped

    padded: list[tuple[int, str, str]] = []
    missing: list[int] = []
    for no, row in data:
        if no in dropped:
            continue
        raw = row[key_idx] if key_idx < len(row) else None
        if values.is_blank(raw):
            missing.append(no)
        elif isinstance(raw, str) and raw != values.normalize_text(raw):
            padded.append((no, raw, values.cell_text(raw)))

    for no, raw, cleaned in padded:
        issues.append(
            Issue(
                kind="key_whitespace",
                severity=rules.severity("key_whitespace"),
                rows=(no,),
                cols=(key_idx + 1,),
                detail=rules.key_whitespace_detail(raw, cleaned),
            )
        )

    if missing:
        dropped.update(missing)
        issues.append(
            Issue(
                kind="blank_key",
                severity=rules.severity("blank_key"),
                rows=tuple(missing),
                detail=rules.blank_key_detail(key_name),
            )
        )

    return dropped


def _duplicate_issues(
    data: list[tuple[int, list[Any]]],
    key_idx: int | None,
    dropped: set[int],
    issues: list[Issue],
) -> None:
    """「整行重复」先于「主键重复」。

    两处都报的话，用户会以为有两条问题行，回原表里只找得到一条。
    一行既是重复行又撞主键时，它归重复行 —— 那条是逐格可验证的，
    用户一眼就能确认自己确实导了两遍。
    """
    alive = [(no, row) for no, row in data if no not in dropped]

    seen_row: dict[tuple[str, ...], int] = {}
    survivors: list[tuple[int, list[Any]]] = []
    for no, row in alive:
        sig = values.row_signature(row)
        first = seen_row.get(sig)
        if first is not None:
            issues.append(
                Issue(
                    kind="duplicate_row",
                    severity=rules.severity("duplicate_row"),
                    rows=(no,),
                    detail=rules.duplicate_row_detail(first),
                )
            )
            continue
        seen_row[sig] = no
        survivors.append((no, row))

    if key_idx is None:
        return

    seen_key: dict[str, int] = {}
    for no, row in survivors:
        raw = row[key_idx] if key_idx < len(row) else None
        key = values.cell_text(raw)
        first = seen_key.get(key)
        if first is not None:
            issues.append(
                Issue(
                    kind="duplicate_key",
                    severity=rules.severity("duplicate_key"),
                    rows=(no,),
                    detail=rules.duplicate_key_detail(key, first),
                )
            )
            continue
        seen_key[key] = no


# ── 类型体检 ──────────────────────────────────────────────────────────
# 「这一列看着是日期/金额，但这几格认不出来」—— 判定用的是模板层同一套解析器
# （`template.is_date` / `is_amount`），所以这里说没问题的值，跑起来一定不会失败。

#: 一列要够长才敢下「这是一列日期」的结论。样本少的时候认错比不认更烦人。
_MIN_VALUES = 3
#: 能解析的至少这么多个 —— 只有一格像日期，那是巧合，不是列的类型
_MIN_PARSED = 2
#: 试解析的顺序。一列要么是日期列要么是金额列，命中一个就停
_TYPED_KINDS = (("日期", template.is_date), ("金额", template.is_amount))


def _unparsable_issues(grid: list[list[Any]], header_row: int, columns: list[Any], issues: list[Issue]) -> None:
    """一列一列地试解析。只有「看着是日期/金额」的列才会往下报。

    判定刻意保守（样本够长 + 严格多数能解析 + 至少 2 格能解析）。
    放宽任何一条，备注列里偶然出现的一个日期串就会把整列打成「日期列」，
    报告全是误报之后，用户就不看体检了。

    一列报一条，`rows` 里放全部坏值：一列 5 处坏值就报 5 条，会把真正要看的东西淹掉。

    只看值本身，与被删掉哪一行无关 —— 所以扫的是**全部**数据行，不是「活下来的行」，
    用户拿到的行号始终能回原文件里对上。
    """
    for col_idx, name in enumerate(columns):
        if values.is_blank(name):
            continue  # 没列名的列不试：模板根本引用不到它（见 blank_header）

        cells: list[tuple[int, Any]] = []
        for no in range(header_row + 1, len(grid) + 1):
            row = grid[no - 1]
            value = row[col_idx] if col_idx < len(row) else None
            if not values.is_blank(value):
                cells.append((no, value))

        if len(cells) < _MIN_VALUES:
            continue

        for what, probe in _TYPED_KINDS:
            flags = [probe(value) for _, value in cells]
            parsed = sum(flags)
            # 严格多数：一半一半等于抛硬币，不足以下结论
            if parsed < _MIN_PARSED or parsed * 2 <= len(cells):
                continue

            bad = [(no, value) for (no, value), ok in zip(cells, flags) if not ok]
            if not bad:
                continue

            rows = tuple(no for no, _ in bad)
            issues.append(
                Issue(
                    kind="unparsable_value",
                    severity=rules.severity("unparsable_value"),
                    rows=rows,
                    cols=(col_idx + 1,),
                    detail=rules.unparsable_value_detail(what, rows, values.cell_text(bad[0][1])),
                )
            )
            break  # 这一列已经定性了，别再拿另一种类型问一遍
