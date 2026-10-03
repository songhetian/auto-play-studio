"""接口层：Excel 对比的完整流程（上传 A/B → 智能匹配 → 执行 → 出报告）。

单测各模块都对，不代表串起来能跑 —— 路由冲突、字段没落库这类问题只有走一遍接口才看得见。
"""
from __future__ import annotations

import os

from openpyxl import load_workbook

from engine import db

CMP_DEFAULT = {
    "tool": "cmp",
    "cmp": {
        "primaryFile": "",
        "primaryColumns": [],
        "primaryFields": [],
        "otherFiles": [],
        "tables": {},
        "maps": {},
        "tolerance": 0.0,
        "outFile": "",
    },
}

A_ROWS = [["A001", 100.0, "张三"], ["A002", 200.5, "李四"], ["A004", 50.0, "赵六"]]
B_ROWS = [["A001", 100.0, "张三"], ["A002", 250.5, "李四"], ["A005", 999.0, "钱七"]]


def upload(client, iid: str, role: str, path: str):
    with open(path, "rb") as f:
        return client.post(
            f"/api/instances/{iid}/compare/upload?role={role}",
            files={"file": (os.path.basename(path), f, "application/vnd.ms-excel")},
        )


def set_field_roles(iid: str, roles: dict[str, str], types: dict[str, str] | None = None):
    cfg = db.get_config(iid)
    types = types or {}
    cfg["cmp"]["primaryFields"] = [
        {"name": f["name"], "role": roles.get(f["name"], "compare"), "type": types.get(f["name"], "text")}
        for f in cfg["cmp"]["primaryFields"]
    ]
    db.save_config(iid, cfg)


def test_full_compare_flow(client, make_instance, xlsx_factory):
    a_path = xlsx_factory(["订单号", "退差金额", "客户"], A_ROWS, name="前台.xlsx")
    b_path = xlsx_factory(["订单编号", "退款差额", "客户"], B_ROWS, name="后台.xlsx")
    iid = make_instance("cmp", CMP_DEFAULT)

    r = upload(client, iid, "primary", a_path)
    assert r.status_code == 200, r.text
    assert r.json()["columns"] == ["订单号", "退差金额", "客户"]
    assert r.json()["rows"] == 3

    # 列名对不上，「订单号」↔「订单编号」、「退差金额」↔「退款差额」全靠智能匹配
    set_field_roles(iid, {"订单号": "key"}, {"退差金额": "number"})

    assert upload(client, iid, "other", b_path).status_code == 200
    mapped = client.post(f"/api/instances/{iid}/compare/auto-map").json()["maps"]
    assert mapped["后台.xlsx"]["订单号"] == "订单编号"
    assert mapped["后台.xlsx"]["退差金额"] == "退款差额"

    res = client.post(f"/api/instances/{iid}/compare/run")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["report"]["summary"] == {"ok": 1, "diff": 1, "missing": 1, "extra": 1, "error": 0, "total": 4}
    assert os.path.exists(body["outFile"])


def test_report_file_is_written_with_two_sheets(client, make_instance, xlsx_factory):
    a_path = xlsx_factory(["订单号", "退差金额"], [["A001", 100.0], ["A002", 200.5]], name="A.xlsx")
    b_path = xlsx_factory(["订单号", "退差金额"], [["A001", 100.0], ["A002", 250.5]], name="B.xlsx")
    iid = make_instance("cmp", CMP_DEFAULT)
    upload(client, iid, "primary", a_path)
    set_field_roles(iid, {"订单号": "key"}, {"退差金额": "number"})
    upload(client, iid, "other", b_path)
    client.post(f"/api/instances/{iid}/compare/auto-map")

    out = client.post(f"/api/instances/{iid}/compare/run").json()["outFile"]

    wb = load_workbook(out)
    try:
        assert wb.sheetnames == ["汇总", "对比结果"]
        assert wb["对比结果"].max_row == 3  # 表头 + 2 行
    finally:
        wb.close()


def test_result_can_be_fetched_after_reload(client, make_instance, xlsx_factory):
    a_path = xlsx_factory(["订单号", "退差金额"], [["A001", 100.0]], name="A.xlsx")
    b_path = xlsx_factory(["订单号", "退差金额"], [["A001", 130.0]], name="B.xlsx")
    iid = make_instance("cmp", CMP_DEFAULT)
    upload(client, iid, "primary", a_path)
    set_field_roles(iid, {"订单号": "key"}, {"退差金额": "number"})
    upload(client, iid, "other", b_path)
    client.post(f"/api/instances/{iid}/compare/auto-map")
    client.post(f"/api/instances/{iid}/compare/run")

    again = client.get(f"/api/instances/{iid}/compare/result")

    assert again.status_code == 200
    assert again.json()["summary"]["diff"] == 1


def test_run_without_primary_file_is_rejected(client, make_instance):
    iid = make_instance("cmp", CMP_DEFAULT)

    r = client.post(f"/api/instances/{iid}/compare/run")

    assert r.status_code == 400
    assert "A 表" in r.json()["detail"]


def test_upload_rejects_non_xlsx(client, make_instance, tmp_path):
    iid = make_instance("cmp", CMP_DEFAULT)
    bad = tmp_path / "a.txt"
    bad.write_text("x")

    with open(bad, "rb") as f:
        r = client.post(
            f"/api/instances/{iid}/compare/upload?role=primary",
            files={"file": ("a.txt", f, "text/plain")},
        )

    assert r.status_code == 400
    assert ".xlsx" in r.json()["detail"]


def test_removing_a_table_clears_its_mapping(client, make_instance, xlsx_factory):
    a_path = xlsx_factory(["订单号", "退差金额"], [["A001", 100.0]], name="A.xlsx")
    b_path = xlsx_factory(["订单号", "退差金额"], [["A001", 100.0]], name="B.xlsx")
    iid = make_instance("cmp", CMP_DEFAULT)
    upload(client, iid, "primary", a_path)
    set_field_roles(iid, {"订单号": "key"}, {"退差金额": "number"})
    upload(client, iid, "other", b_path)
    client.post(f"/api/instances/{iid}/compare/auto-map")

    assert client.delete(f"/api/instances/{iid}/compare/table?name=B.xlsx").status_code == 200

    cfg = db.get_config(iid)["cmp"]
    assert cfg["tables"] == {}
    assert cfg["maps"] == {}


def test_compare_tool_runs_through_the_instance_runner(client, make_instance, xlsx_factory):
    """从「运行详情」页点开始执行也要能跑通（走 runner，不是直连接口）。"""
    a_path = xlsx_factory(["订单号", "退差金额"], [["A001", 100.0], ["A002", 200.0]], name="A.xlsx")
    b_path = xlsx_factory(["订单号", "退差金额"], [["A001", 100.0], ["A002", 260.0]], name="B.xlsx")
    iid = make_instance("cmp", CMP_DEFAULT)
    upload(client, iid, "primary", a_path)
    set_field_roles(iid, {"订单号": "key"}, {"退差金额": "number"})
    upload(client, iid, "other", b_path)
    client.post(f"/api/instances/{iid}/compare/auto-map")

    assert client.post(f"/api/instances/{iid}/control/start").status_code == 200

    from tests.conftest import _wait_status

    assert _wait_status(iid, {"completed"}, 10) == "completed"
    rows = client.get(f"/api/instances/{iid}/rows").json()
    assert [r["status"] for r in rows] == ["ok", "diff"]
