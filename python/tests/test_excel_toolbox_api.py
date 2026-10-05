"""Excel 工具箱的接口层测试：真文件进、真文件出。

覆盖"动态列"这个核心承诺 —— 传几列、传什么列名，产出的表就是那个形状。
"""
from __future__ import annotations

import csv
import pathlib

import pytest


@pytest.fixture
def main_xlsx(xlsx_factory):
    """主表：订单号 / 买家 / 金额"""
    return xlsx_factory(
        ["订单号", "买家", "金额"],
        [["SO1", "张", 10], ["SO2", "李", 20], ["SO9", "无", 30]],
        name="订单表.xlsx",
    )


@pytest.fixture
def src_xlsx(xlsx_factory):
    """来源表：订单/客户/物流状态/快递单"""
    return xlsx_factory(
        ["订单", "客户", "物流状态", "快递单"],
        [["SO1", "张", "已签收", "SF1"], ["SO2", "李", "运输中", "YT2"]],
        name="物流表.xlsx",
    )


def read_xlsx(path: str) -> list[dict]:
    """读回产出的 xlsx。

    把 None 归一成 ""：xlsx 分不出「空串」和「空单元格」，
    所以断言按"这里是空的"来写，而不是钉住 None 这个存储细节。
    """
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.values)
    if not rows:
        return []
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, ("" if v is None else v for v in r))) for r in rows[1:]]


def read_headers(path: str) -> list[str]:
    from openpyxl import load_workbook

    ws = load_workbook(path).active
    return [str(c) for c in next(ws.values)]


class TestColumns:
    """只读列名：用户还没选列时，页面也得能显示可选的列。"""

    def test_返回列名与行数(self, client, main_xlsx):
        r = client.get("/api/excel-toolbox/columns", params={"path": main_xlsx})
        assert r.status_code == 200
        body = r.json()
        assert body["columns"] == ["订单号", "买家", "金额"]
        assert body["rowCount"] == 3

    def test_文件不存在给404而不是500(self, client):
        r = client.get("/api/excel-toolbox/columns", params={"path": "E:/nope.xlsx"})
        assert r.status_code == 404


class TestMerge:
    def test_多键匹配补全并产出新文件(self, client, main_xlsx, src_xlsx):
        r = client.post(
            "/api/excel-toolbox/merge",
            json={
                "mainPath": main_xlsx,
                "srcPath": src_xlsx,
                "mainKeys": ["订单号", "买家"],
                "srcKeys": ["订单", "客户"],
                "fillColumns": ["物流状态", "快递单"],
            },
        )
        assert r.status_code == 200
        body = r.json()
        out = read_xlsx(body["path"])
        assert [x["物流状态"] for x in out] == ["已签收", "运输中", ""]
        # 原文件只读，产出是新文件
        assert body["path"] != main_xlsx
        assert body["rowCount"] == 3
        assert body["matchedCount"] == 2

    def test_补全列可改名(self, client, main_xlsx, src_xlsx):
        r = client.post(
            "/api/excel-toolbox/merge",
            json={
                "mainPath": main_xlsx,
                "srcPath": src_xlsx,
                "mainKeys": ["订单号"],
                "srcKeys": ["订单"],
                "fillColumns": [{"from": "快递单", "to": "快递单号"}],
            },
        )
        assert read_xlsx(r.json()["path"])[0]["快递单号"] == "SF1"

    def test_不给补全列时也能跑_只做键校验(self, client, main_xlsx, src_xlsx):
        r = client.post(
            "/api/excel-toolbox/merge",
            json={
                "mainPath": main_xlsx,
                "srcPath": src_xlsx,
                "mainKeys": ["订单号"],
                "srcKeys": ["订单"],
                "fillColumns": [],
            },
        )
        assert r.status_code == 200
        assert r.json()["rowCount"] == 3

    def test_键列不存在给400_并说清哪一列(self, client, main_xlsx, src_xlsx):
        r = client.post(
            "/api/excel-toolbox/merge",
            json={
                "mainPath": main_xlsx,
                "srcPath": src_xlsx,
                "mainKeys": ["不存在的列"],
                "srcKeys": ["订单"],
                "fillColumns": ["物流状态"],
            },
        )
        assert r.status_code == 400
        assert "不存在的列" in r.json()["detail"]


class TestGenerate:
    def test_按勾选列生成新表_列序跟着勾选走(self, client, main_xlsx):
        r = client.post(
            "/api/excel-toolbox/generate",
            json={"path": main_xlsx, "columns": ["金额", "订单号"]},
        )
        assert r.status_code == 200
        assert read_headers(r.json()["path"]) == ["金额", "订单号"]
        assert r.json()["rowCount"] == 3

    def test_空列也接受_产出只有表头(self, client, main_xlsx):
        r = client.post("/api/excel-toolbox/generate", json={"path": main_xlsx, "columns": []})
        assert r.status_code == 200
        assert read_xlsx(r.json()["path"]) == []


class TestDedupe:
    def test_按业务键去重并报去掉了多少(self, client, xlsx_factory):
        p = xlsx_factory(["订单号", "金额"], [["A", 1], ["A", 2], ["B", 3]], name="dup.xlsx")
        r = client.post("/api/excel-toolbox/dedupe", json={"path": p, "keys": ["订单号"]})
        body = r.json()
        assert body["rowCount"] == 2
        assert body["removedCount"] == 1

    def test_不传键时按整行去重(self, client, xlsx_factory):
        p = xlsx_factory(["订单号", "金额"], [["A", 1], ["A", 1], ["B", 3]], name="dup2.xlsx")
        assert client.post("/api/excel-toolbox/dedupe", json={"path": p, "keys": []}).json()["removedCount"] == 1


class TestConvert:
    def test_xlsx转csv(self, client, main_xlsx):
        r = client.post("/api/excel-toolbox/convert", json={"path": main_xlsx, "target": "csv"})
        assert r.status_code == 200
        with open(r.json()["path"], newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
        assert rows[0] == ["订单号", "买家", "金额"]
        assert len(rows) == 4

    def test_csv转xlsx(self, client, tmp_path):
        p = tmp_path / "in.csv"
        p.write_text("订单号,金额\nA1,5\n", encoding="utf-8")
        r = client.post("/api/excel-toolbox/convert", json={"path": str(p), "target": "xlsx"})
        assert read_headers(r.json()["path"]) == ["订单号", "金额"]

    def test_不支持的目标格式给400(self, client, main_xlsx):
        r = client.post("/api/excel-toolbox/convert", json={"path": main_xlsx, "target": "pdf"})
        assert r.status_code == 400


class TestSplit:
    def test_按列拆成多个文件(self, client, xlsx_factory):
        p = xlsx_factory(["店铺", "金额"], [["A店", 1], ["B店", 2], ["A店", 3]], name="shops.xlsx")
        r = client.post("/api/excel-toolbox/split", json={"path": p, "column": "店铺"})
        body = r.json()
        assert len(body["files"]) == 2
        assert sorted(f["value"] for f in body["files"]) == ["A店", "B店"]
        # 每个文件真的落盘了
        assert all(pathlib.Path(f["path"]).exists() for f in body["files"])

    def test_列不存在给400(self, client, main_xlsx):
        assert client.post("/api/excel-toolbox/split", json={"path": main_xlsx, "column": "无"}).status_code == 400


class TestFilterExport:
    def test_按列等值筛出子集(self, client, xlsx_factory):
        p = xlsx_factory(["状态", "金额"], [["已签收", 1], ["运输中", 2]], name="st.xlsx")
        r = client.post(
            "/api/excel-toolbox/filter",
            json={"path": p, "column": "状态", "value": "已签收"},
        )
        assert r.status_code == 200
        assert r.json()["rowCount"] == 1

    def test_空值返回全部(self, client, main_xlsx):
        r = client.post("/api/excel-toolbox/filter", json={"path": main_xlsx, "column": "买家", "value": ""})
        assert r.json()["rowCount"] == 3


class TestProblemExport:
    def test_导出行号与原因_便于回查原表(self, client, xlsx_factory):
        # 第 2 行主键为空、第 4 行主键与第 3 行重复 —— 两类问题各一行
        p = xlsx_factory(["订单号", "金额"], [[None, 1], ["B", 2], ["B", 3]], name="bad.xlsx")
        r = client.post(
            "/api/excel-toolbox/problem-rows",
            json={"path": p, "keyColumn": "订单号"},
        )
        assert r.status_code == 200
        body = r.json()
        # 问题行要带上 Excel 行号，方便直接定位回去改
        assert "行号" in read_headers(body["path"])
        rows = read_xlsx(body["path"])
        # 空主键 + 主键重复，两行都该被认出来
        assert body["rowCount"] == 2
        assert any("主键「订单号」为空" in str(x["问题"]) for x in rows)
        assert any("重复" in str(x["问题"]) for x in rows)
        # 行号是 Excel 真实行号（表头占第 1 行）
        assert {x["行号"] for x in rows} == {2, 4}


class TestCsvSupport:
    """界面各面板都标注「xlsx / csv」，引擎就必须真的都认 csv。

    曾经 columns / convert 支持 csv，而 merge / dedupe / filter 等走 read_table
    的入口一律 400 —— 用户按提示传 csv 会被打断。
    """

    def test_filter_能读csv(self, client, tmp_path):
        p = tmp_path / "订单.csv"
        p.write_text("订单号,状态\nSO1,已签收\nSO2,运输中\n", encoding="utf-8-sig")

        r = client.post(
            "/api/excel-toolbox/filter",
            json={"path": str(p), "column": "状态", "value": "已签收"},
        )

        assert r.status_code == 200, r.text
        assert r.json()["rowCount"] == 1
