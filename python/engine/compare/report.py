"""对比报告写出：汇总 + 对比结果两张表，带配色。

只负责「把 compare_core 的结果变成一份能发出去的 xlsx」，
读表在 excel_io.py，比对在 compare_core.py —— 三者互不依赖，可单独替换。

输入结构见 compare_core.compare()：
    {header, rows:[{cells,status,cell_status,...}], summary, warnings, cmp_fields}
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

#: 结论 -> 单元格底色
STATUS_FILL = {
    "ok": "C6EFCE",       # 绿：一致
    "diff": "FFC7CE",     # 红：不一致
    "missing": "BDD7EE",  # 蓝：缺失（A 表有，对比表没有）
    "extra": "E4DFEC",    # 紫：多余（仅对比表有）
    "error": "F8CBAD",    # 橙：数据错误（该是数字却不是数字）
}

_SUMMARY_LABELS = [
    ("一致", "ok"),
    ("不一致", "diff"),
    ("缺失(A表独有)", "missing"),
    ("多余(仅对比表有)", "extra"),
    ("数据错误", "error"),
    ("合计", None),
]

_HEADER_FILL = PatternFill("solid", fgColor="305496")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_BORDER = Border(*(Side(style="thin", color="D9D9D9") for _ in range(4)))
_WRAP = Alignment(vertical="center", wrap_text=True)


def _col_letter(idx: int) -> str:
    """1-based 列号 -> 列字母（超过 Z 也能正确进位）。"""
    name = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        name = chr(ord("A") + rem) + name
    return name


def write_report(path: str, report: dict, primary_name: str, tolerance: float) -> str:
    """写出报告，返回文件路径。生成『汇总』与『对比结果』两个工作表。"""
    summary = report.get("summary", {})
    wb = Workbook()
    wb.remove(wb.active)

    # ── 汇总 ──
    sm = wb.create_sheet("汇总")
    sm.append(["Excel 多表对比报告（以 A 表为基准）"])
    sm["A1"].font = Font(bold=True, size=14)
    sm.append(["主表", primary_name])
    sm.append(["数值容差", tolerance])
    sm.append([])
    sm.append(["结论", "行数"])
    for c in (1, 2):
        cell = sm.cell(row=5, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for label, key in _SUMMARY_LABELS:
        sm.append([label, summary.get(key if key else "total", 0)])
        if key and key in STATUS_FILL:
            sm.cell(row=sm.max_row, column=1).fill = PatternFill("solid", fgColor=STATUS_FILL[key])
    if report.get("warnings"):
        sm.append([])
        sm.append(["提示/告警"])
        sm.cell(row=sm.max_row, column=1).font = Font(bold=True)
        for w in report["warnings"]:
            sm.append([w])
    sm.column_dimensions["A"].width = 36
    sm.column_dimensions["B"].width = 12

    # ── 对比结果 ──
    ws = wb.create_sheet("对比结果")
    header = report["header"]
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER

    for row in report["rows"]:
        ws.append([row["cells"].get(h, "") for h in header])
        r = ws.max_row
        status = row.get("status")
        if status in STATUS_FILL:
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
        for h, cell_status in row.get("cell_status", {}).items():
            if cell_status in STATUS_FILL and h in header:
                ws.cell(row=r, column=header.index(h) + 1).fill = PatternFill(
                    "solid", fgColor=STATUS_FILL[cell_status]
                )
        for c in range(1, len(header) + 1):
            ws.cell(row=r, column=c).border = _BORDER
            ws.cell(row=r, column=c).alignment = _WRAP

    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 12
    for i in range(2, len(header) + 1):
        ws.column_dimensions[_col_letter(i)].width = 16

    wb.save(path)
    return path
