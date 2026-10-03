"""Seam 1：Excel 读写与状态回写。

以真实 xlsx 文件为真相来源，断言「执行一次后文件长什么样、下次读出来是什么」。
"""
from __future__ import annotations

import os

from openpyxl import load_workbook

from engine import excel


def _row_text(path: str, row_no: int) -> str:
    """把某一行的所有单元格拼成文本，用于断言「原因写进了这一行」而不耦合具体列。"""
    wb = load_workbook(path)
    try:
        ws = wb[wb.sheetnames[0]]
        return " ".join(str(c) for c in next(ws.iter_rows(min_row=row_no, max_row=row_no, values_only=True)) if c)
    finally:
        wb.close()


def test_detect_columns(order_xlsx):
    assert excel.detect_columns(order_xlsx) == ["订单编号", "客户名称", "话术模板"]


def test_status_column_is_created_once(order_xlsx):
    first = excel.ensure_status_column(order_xlsx)
    assert excel.detect_columns(order_xlsx)[-1] == excel.STATUS_COLUMN_DEFAULT

    second = excel.ensure_status_column(order_xlsx)
    assert second == first
    assert excel.detect_columns(order_xlsx).count(excel.STATUS_COLUMN_DEFAULT) == 1


def test_success_is_written_and_read_back(order_xlsx):
    excel.ensure_status_column(order_xlsx)
    excel.write_row_status(order_xlsx, 2, excel.STATUS_OK)

    rows = excel.read_rows(order_xlsx, "客户名称", 1, 100)
    assert [(r["row_no"], r["key"]) for r in rows] == [(2, "张三"), (3, "李四")]
    assert rows[0]["existing_status"] == excel.STATUS_OK
    assert rows[1]["existing_status"] == ""


def test_failure_writes_status_and_reason(order_xlsx):
    excel.ensure_status_column(order_xlsx)
    excel.write_row_status(order_xlsx, 2, excel.STATUS_OK)
    excel.write_row_status(order_xlsx, 3, excel.STATUS_ERR, "未找到窗口")

    rows = excel.read_rows(order_xlsx, "客户名称", 1, 100)
    assert rows[1]["existing_status"] == excel.STATUS_ERR
    assert "未找到窗口" in _row_text(order_xlsx, 3)


def test_read_rows_carries_the_whole_row(order_xlsx):
    """指令里要能引用 {客户名称}、消息内容列也要按行取值 —— 光有主键不够用。"""
    rows = excel.read_rows(order_xlsx, "客户名称", 1, 100)

    assert rows[0]["values"] == {"订单编号": "A001", "客户名称": "张三", "话术模板": "您好，新品上架"}
    assert rows[1]["values"]["话术模板"] == "您好，老客回访"


def test_read_rows_does_not_expose_the_status_column_as_data(order_xlsx):
    """状态列是程序自己的记账栏，不该被当成可引用字段。"""
    excel.write_row_status(order_xlsx, 2, excel.STATUS_OK)

    values = excel.read_rows(order_xlsx, "客户名称", 1, 100)[0]["values"]

    assert excel.STATUS_COLUMN_DEFAULT not in values
    assert "客户名称" in values


def test_custom_status_column_is_respected(order_xlsx):
    """用户指定了自己的状态列时，写入与跳过判断都必须走这一列。"""
    excel.ensure_status_column(order_xlsx, "发送结果")
    excel.write_row_status(order_xlsx, 2, excel.STATUS_OK, status_col="发送结果")

    rows = excel.read_rows(order_xlsx, "客户名称", 1, 100, status_col="发送结果")
    assert [r["existing_status"] for r in rows] == [excel.STATUS_OK, ""]


def test_read_rows_skips_blank_key_and_respects_range(xlsx_factory):
    path = xlsx_factory(
        ["客户名称"],
        [["张三"], [""], ["李四"], ["王五"]],
        name="blank.xlsx",
    )
    excel.ensure_status_column(path)
    rows = excel.read_rows(path, "客户名称", 1, 100)
    assert [r["row_no"] for r in rows] == [2, 4, 5]

    assert [r["key"] for r in excel.read_rows(path, "客户名称", 4, 4)] == ["李四"]


def test_backup_file_keeps_a_copy(order_xlsx):
    dst = excel.backup_file(order_xlsx)
    assert os.path.exists(dst)
    assert dst.endswith("_backup.xlsx")
    assert excel.detect_columns(dst) == excel.detect_columns(order_xlsx)


def test_logistics_result_goes_to_a_new_file(xlsx_factory):
    src = xlsx_factory(
        ["订单号", "物流单号"],
        [["O1", "SF1234567890"], ["O2", "UNKNOWN"]],
        name="waybill.xlsx",
    )
    before = excel.detect_columns(src)
    results = [
        {"no": "SF1234567890", "company": "顺丰速运", "status": "已签收", "signed_at": "10-01 15:22", "trace": "已签收"},
    ]

    out = excel.write_logistics_result(src, results, "物流单号")

    assert out != src
    assert excel.detect_columns(src) == before  # 原文件保持只读
    assert excel.detect_columns(out)[-4:] == ["物流公司", "物流状态", "签收时间", "最新轨迹"]
    assert "顺丰速运" in _row_text(out, 2)
    assert "顺丰速运" not in _row_text(out, 3)  # 未匹配到的行留空
