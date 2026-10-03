"""Seam：报告写出。

产出一份能直接发出去的 xlsx：『汇总』给人看结论，『对比结果』给人对账。
断言方式是把它读回来 —— 真实文件是唯一的真相来源。
"""
from __future__ import annotations

from openpyxl import load_workbook

from engine.compare.compare_core import compare
from engine.compare.report import write_report


def build_report(front_backend, primary_fields, tolerance=0.0):
    a, b = front_backend
    maps = {"订单号": "订单号", "退差金额": "退差金额", "客户": "客户"}
    return compare(a, primary_fields, [{"table": b, "maps": maps}], tolerance)


def _summary_rows(ws):
    return {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}


def test_report_has_two_sheets(tmp_path, front_backend, primary_fields):
    report = build_report(front_backend, primary_fields)
    out = str(tmp_path / "报告.xlsx")

    write_report(out, report, "前台.xlsx", 0.01)

    wb = load_workbook(out)
    assert wb.sheetnames == ["汇总", "对比结果"]
    wb.close()


def test_summary_carries_the_parameters_and_counts(tmp_path, front_backend, primary_fields):
    report = build_report(front_backend, primary_fields)
    out = str(tmp_path / "报告.xlsx")
    write_report(out, report, "前台.xlsx", 0.01)

    wb = load_workbook(out)
    try:
        rows = _summary_rows(wb["汇总"])
    finally:
        wb.close()

    assert rows["主表"] == "前台.xlsx"
    assert rows["数值容差"] == 0.01
    assert rows["一致"] == report["summary"]["ok"] == 2
    assert rows["不一致"] == 1
    assert rows["缺失(A表独有)"] == 1
    assert rows["多余(仅对比表有)"] == 1
    assert rows["数据错误"] == 1
    assert rows["合计"] == 6


def test_result_sheet_matches_the_report(tmp_path, front_backend, primary_fields):
    report = build_report(front_backend, primary_fields)
    out = str(tmp_path / "报告.xlsx")
    write_report(out, report, "前台.xlsx", 0.0)

    wb = load_workbook(out)
    try:
        ws = wb["对比结果"]
        header = [c.value for c in ws[1]]
        first_row = [c.value for c in ws[2]]
    finally:
        wb.close()

    assert header == report["header"]
    assert first_row == [report["rows"][0]["cells"][h] for h in header]
    assert first_row[0] == "一致"


def test_diff_cell_is_highlighted(tmp_path, front_backend, primary_fields):
    report = build_report(front_backend, primary_fields)
    out = str(tmp_path / "报告.xlsx")
    write_report(out, report, "前台.xlsx", 0.0)

    wb = load_workbook(out)
    try:
        ws = wb["对比结果"]
        header = [c.value for c in ws[1]]
        col = header.index("后台.xlsx·退差金额") + 1
        # A002 是唯一不一致的行（表头占第 1 行，A001 在第 2 行）
        diff_row = 3
        fill = ws.cell(row=diff_row, column=col).fill
        conclusion_fill = ws.cell(row=diff_row, column=1).fill
    finally:
        wb.close()

    assert str(fill.fgColor.rgb).endswith("FFC7CE"), "不一致的单元格应为红色"
    assert str(conclusion_fill.fgColor.rgb).endswith("FFC7CE"), "结论列同样标红"


def test_warnings_are_written_into_the_summary(tmp_path, front_backend, primary_fields):
    a, b = front_backend
    b["data"].append({"订单号": "A001", "退差金额": 111.0, "客户": "张三"})
    report = compare(a, primary_fields, [{"table": b, "maps": {"订单号": "订单号", "退差金额": "退差金额"}}], 0.0)

    out = str(tmp_path / "报告.xlsx")
    write_report(out, report, "前台.xlsx", 0.0)

    wb = load_workbook(out)
    try:
        texts = [str(cell) for row in wb["汇总"].iter_rows(values_only=True) for cell in row]
    finally:
        wb.close()

    assert any("重复" in t for t in texts if t)


def test_report_is_a_valid_workbook_that_can_be_reopened(tmp_path, front_backend, primary_fields):
    """写出的文件必须能被 openpyxl 正常读回（冻结窗格、列宽等样式不能破坏文件）。"""
    report = build_report(front_backend, primary_fields)
    out = str(tmp_path / "报告.xlsx")
    write_report(out, report, "前台.xlsx", 0.0)

    wb = load_workbook(out)
    try:
        assert wb["对比结果"].max_row == len(report["rows"]) + 1
    finally:
        wb.close()
