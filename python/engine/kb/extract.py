"""抽取层：把 PDF / Word / Excel / PPT / 纯文本读成可切块的纯文本。

两块约定：

1. **块与块之间用空行分开**（`BLOCK_SEP`）。上层 `chunk_text` 就是按空行分块的，
   所以「一个自然段 / 一张表 / 一页 PDF」在这里各成为一块 —— 命中时片段的粒度
   就等于这里的粒度。
   **一张表算一块，行与行之间用单个换行**：`chunk_text` 的句子边界正则把 `\n` 也当边界，
   所以短表整块保留（片段里能连着看到表头）、长表自动在行边界切开，不需要这里再分。
2. **纯文本剥掉首尾空白**。同一份内容从 md 抽和从 docx 抽应当得到一样的串，
   首尾空行没有信息量。

不剥正文内部的空白：空行是要传给 `chunk_text` 用的。
"""
from __future__ import annotations

import dataclasses
import pathlib

#: 纳入索引的扩展名（与 know 的 `SUPPORTED_EXT` 一致）
SUPPORTED_EXT = ("pdf", "docx", "xlsx", "pptx", "md", "txt")

#: 单个文件的文本上限：一本几 MB 的手册不该把整个索引拖垮
MAX_CHARS = 500_000

#: 纯文本的编码尝试顺序。Windows 上 GBK 老文件和带 BOM 的 UTF-8 都很常见；
#: `latin-1` 兜底永不失败，保证这个循环一定收敛。
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gbk", "latin-1")

#: 表格单元格的连接符。`A001 | 张三 | 退款` 在命中片段里比制表符可读得多。
CELL_SEP = " | "

#: 块之间的分隔（上层按空行分块）
BLOCK_SEP = "\n\n"


class ExtractError(Exception):
    """文件读不出来。

    消息里**必须带文件名** —— 知识库文件夹里有一个坏文件是常态，
    而调用方（索引器）会把它记成「这个文件跳过了」，用户得知道是哪个。
    """


class UnsupportedType(ExtractError):
    """扩展名不在 `SUPPORTED_EXT` 里。

    继承 `ExtractError` 是为了让索引器只有一个 catch 点：对它而言
    「这个类型不支持」与「这个文件坏了」都是「跳过一个文件」，不是两种事故。
    """


@dataclasses.dataclass(frozen=True)
class RawFile:
    """一个已经读出文本的文件。对应 know 渲染层的 `RawFile`。"""

    #: 文件名（不含目录）
    name: str
    #: 绝对路径，原样保存（路径归一化留给比较层做）
    path: str
    #: 文件字节数；**不随文本截断变化**
    size: int
    #: 修改时间，epoch 秒。展示格式交给前端，避免跨进程时区歧义
    mtime: float
    #: 抽出的文本，已按 `MAX_CHARS` 截断
    text: str
    #: 是否发生了截断
    truncated: bool


def ext_of(path: str | pathlib.Path) -> str:
    """小写扩展名；没有扩展名返回空串。"""
    name = pathlib.Path(path).name
    i = name.rfind(".")
    return name[i + 1 :].lower() if i >= 0 else ""


def is_supported(path: str | pathlib.Path) -> bool:
    return ext_of(path) in SUPPORTED_EXT


def extract_text(path: str | pathlib.Path) -> str:
    """按扩展名分派抽取，返回剥掉首尾空白的文本。

    不支持的扩展名抛 `UnsupportedType`（在读文件之前就抛，不去碰内容）；
    解析失败抛 `ExtractError`，消息带文件名与原始异常，便于用户自己去看那个文件。
    """
    p = pathlib.Path(path)
    ext = ext_of(p)
    if ext not in SUPPORTED_EXT:
        raise UnsupportedType(f"{p.name}：{ext or '无扩展名'} 不在可索引的类型里（{'/'.join(SUPPORTED_EXT)}）")

    try:
        text = _EXTRACTORS[ext](p)
    except Exception as e:  # noqa: BLE001 —— 任何解析异常都要变成带文件名的可读报错
        raise ExtractError(f"{p.name}：{type(e).__name__}: {e}") from e
    # Windows 上的 txt / 从 Word 复制来的段落普遍带 CRLF。留着 `\r` 会让片段显示多出
    # 看不见的字符，也会让高亮偏移算错 —— 统一成 LF，只在这一处做。
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def read_raw_file(path: str | pathlib.Path) -> RawFile:
    """读取一个文件：元信息 + 文本 + 截断标记。"""
    p = pathlib.Path(path)
    st = p.stat()  # 先 stat：文件不存在时直接抛 FileNotFoundError，比解析到一半再炸清楚
    text = extract_text(p)
    truncated = len(text) > MAX_CHARS
    if truncated:
        text = text[:MAX_CHARS]
    return RawFile(name=p.name, path=str(p), size=st.st_size, mtime=st.st_mtime, text=text, truncated=truncated)


# ── 各格式解析 ────────────────────────────────────────────────────────


def _short_text_path(path: str | pathlib.Path) -> str:
    """md / txt 原样读，按 `TEXT_ENCODINGS` 逐个试。"""
    raw = pathlib.Path(path).read_bytes()
    for enc in TEXT_ENCODINGS:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ExtractError(f"{pathlib.Path(path).name}：试过 {'/'.join(TEXT_ENCODINGS)} 都解不出来")


def _docx(path: str | pathlib.Path) -> str:
    """段落与表格**按文档顺序**读。

    不用 `doc.paragraphs` + `doc.tables`（那是两个独立集合，会丢掉相对顺序），
    而是直接遍历 body 的子元素。
    """
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, doc).text.strip()
        elif tag == "tbl":
            t = "\n".join(_table_lines(Table(child, doc)))
        else:
            continue
        if t:
            blocks.append(t)
    return BLOCK_SEP.join(blocks)


def _table_lines(table) -> list[str]:
    """表格 → 行文本列表（单元格按 `CELL_SEP` 连）。整行为空的跳过。"""
    lines = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        if any(cells):
            lines.append(CELL_SEP.join(cells))
    return lines


def _xlsx(path: str | pathlib.Path) -> str:
    """每个非空工作表一块，行与行用换行分隔（与 docx / pptx 的表格同一套规则）。

    不加「工作表：Sheet1」这样的标记 —— 命中片段要干净，是哪个文件已经由结果本身给出。
    代价是同一工作簿里不同工作表的同名行看不出区别，这个代价可以接受。
    """
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    sheets: list[str] = []
    try:
        for ws in wb.worksheets:
            lines = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    lines.append(CELL_SEP.join(c for c in cells if c))
            if lines:
                sheets.append("\n".join(lines))
    finally:
        wb.close()
    return BLOCK_SEP.join(sheets)


def _pptx(path: str | pathlib.Path) -> str:
    """按幻灯片顺序，每个文本框一块、每张表一块。"""
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    blocks.append(t)
            elif shape.has_table:
                lines = _table_lines(shape.table)
                if lines:
                    blocks.append("\n".join(lines))
    return BLOCK_SEP.join(blocks)


def _pdf(path: str | pathlib.Path) -> str:
    """每页一块。

    能抽出的只有 PDF 里真实存在的文本层 —— 扫描件（整页图片）抽出来是空的，
    这属于 OCR（待办 #10）的范围，这里不假装能做到。
    """
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return BLOCK_SEP.join(t for t in pages if t)


_EXTRACTORS = {
    "md": _short_text_path,
    "txt": _short_text_path,
    "docx": _docx,
    "xlsx": _xlsx,
    "pptx": _pptx,
    "pdf": _pdf,
}
