"""Excel 读写：列检测、状态回写（成功/失败/原因）、跳过成功行、物流结果写入新文件。"""
from __future__ import annotations

import os
import re
import shutil
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font

STATUS_COLUMN_DEFAULT = "执行状态"


def normalize_no(value: Any) -> str:
    """单号归一化：去掉首尾与中间的空白（含全角空格）。

    从平台导出或网页复制的单号常带空格，不归一化的话永远匹配不上。
    """
    if value is None:
        return ""
    return re.sub(r"[\s\u3000]+", "", str(value))
STATUS_OK = "成功"
STATUS_ERR = "失败"
STATUS_SKIP = "已跳过"


def detect_columns(path: str) -> list[str]:
    """读取首个工作表的表头列名。"""
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return [str(h) for h in headers if h is not None]
    finally:
        wb.close()


def read_rows(
    path: str,
    key_col: str,
    from_row: int,
    to_row: int,
    status_col: str = STATUS_COLUMN_DEFAULT,
) -> list[dict[str, Any]]:
    """按列名读取指定行区间的数据（行号与 Excel 行号一致）。

    existing_status 取自 status_col —— 「跳过已成功的行」完全依赖它，
    因此用户指定了自定义状态列时必须读那一列，否则跳过逻辑会失效。

    values 是这一行的完整数据（不含状态列），供指令里的 {列名} 占位符与
    「消息内容列」取值 —— 只有主键的话，每行发不同内容是做不到的。
    """
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = [str(h) if h is not None else "" for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        if key_col not in headers:
            raise ValueError(f"未找到列「{key_col}」，可用列：{headers}")
        key_idx = headers.index(key_col)
        status_idx = headers.index(status_col) if status_col in headers else None
        out: list[dict[str, Any]] = []
        for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if r_idx < from_row or r_idx > to_row:
                continue
            value = row[key_idx] if key_idx < len(row) else None
            if value is None or str(value).strip() == "":
                continue
            raw = row[status_idx] if status_idx is not None and status_idx < len(row) else None
            values = {
                name: (row[i] if i < len(row) else None)
                for i, name in enumerate(headers)
                if name and name != status_col
            }
            out.append(
                {
                    "row_no": r_idx,
                    "key": str(value),
                    "existing_status": "" if raw is None else str(raw).strip(),
                    "values": values,
                }
            )
        return out
    finally:
        wb.close()


def sample_row(path: str, key_col: str = "", status_col: str = STATUS_COLUMN_DEFAULT) -> dict[str, Any]:
    """第一行有内容的数据，给配置页的实时预览当样例。

    跳过前导空行：表头下面空一行很常见，取第 2 行就会预览出一片空白，等于没预览。
    values 里保留原始类型（数字是数字、日期是 datetime），
    只有拿真值预览，用户才看得到 `{金额|money}` 到底会输出成什么样。
    """
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = [str(h) if h is not None else "" for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            values = {
                name: (row[i] if i < len(row) else None)
                for i, name in enumerate(headers)
                if name and name != status_col
            }
            if not any(v is not None and str(v).strip() != "" for v in values.values()):
                continue
            key_idx = headers.index(key_col) if key_col in headers else None
            raw = row[key_idx] if key_idx is not None and key_idx < len(row) else None
            return {"row_no": r_idx, "key": "" if raw is None else str(raw).strip(), "values": values}
        return {"row_no": 2, "key": "", "values": {}}
    finally:
        wb.close()


def count_rows(path: str) -> int:
    """数据行数（不含表头）—— 上传后用来给出默认行区间。"""
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        return sum(1 for _ in ws.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()


def ensure_status_column(path: str, status_col: str = STATUS_COLUMN_DEFAULT) -> int:
    """状态列不存在时自动创建，返回该列索引（1-based）。"""
    wb = load_workbook(path)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = [c.value for c in ws[1]]
        if status_col in headers:
            return headers.index(status_col) + 1
        col = ws.max_column + 1
        ws.cell(row=1, column=col, value=status_col).font = Font(bold=True)
        wb.save(path)
        return col
    finally:
        wb.close()


def write_row_status(path: str, row_no: int, status: str, reason: str = "", status_col: str = STATUS_COLUMN_DEFAULT) -> None:
    """把执行结果写回 Excel 状态列：成功 / 失败 / 失败原因。"""
    wb = load_workbook(path)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = [c.value for c in ws[1]]
        if status_col not in headers:
            ensure_status_column(path, status_col)
            wb.close()
            wb = load_workbook(path)
            ws = wb[wb.sheetnames[0]]
            headers = [c.value for c in ws[1]]
        col = headers.index(status_col) + 1
        ws.cell(row=row_no, column=col, value=status)
        if reason:
            reason_col = headers.index("失败原因") + 1 if "失败原因" in headers else ws.max_column + 1
            if "失败原因" not in headers:
                ws.cell(row=1, column=reason_col, value="失败原因").font = Font(bold=True)
            ws.cell(row=row_no, column=reason_col, value=reason)
        wb.save(path)
    finally:
        wb.close()


def backup_file(path: str) -> str:
    dst = re.sub(r"(\.[^.]+)$", r"_backup\1", path)
    shutil.copy2(path, dst)
    return dst


def write_logistics_result(src_path: str, results: list[dict[str, str]], waybill_col: str) -> str:
    """物流查询：原文件保持只读，结果写入新文件（原名_物流信息.xlsx）。"""
    out_path = re.sub(r"(\.xlsx|\.xls|\.csv)$", r"_物流信息\1", src_path, flags=re.I)
    if not out_path.endswith(".xlsx"):
        out_path += ".xlsx"
    shutil.copy2(src_path, out_path)

    wb = load_workbook(out_path)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = [c.value for c in ws[1]]
        base = ws.max_column
        new_cols = ["物流公司", "物流状态", "签收时间", "最新轨迹"]
        for i, name in enumerate(new_cols, start=1):
            ws.cell(row=1, column=base + i, value=name).font = Font(bold=True)

        if waybill_col not in headers:
            raise ValueError(f"未找到物流单号列「{waybill_col}」")
        no_idx = headers.index(waybill_col)

        by_no = {normalize_no(r.get("no", "")): r for r in results}
        for r_idx in range(2, ws.max_row + 1):
            no = ws.cell(row=r_idx, column=no_idx + 1).value
            r = by_no.get(normalize_no(no)) if no is not None else None
            if not r:
                continue
            for i, key in enumerate(["company", "status", "signed_at", "trace"], start=1):
                ws.cell(row=r_idx, column=base + i, value=r.get(key, ""))
        wb.save(out_path)
    finally:
        wb.close()
    return out_path
