"""整理：照着体检的结论出一份新文件，并留下逐处改动台账。原文件一个字节都不动。

**整理不自己判一遍。** 它先要一份 `inspector` 的报告，然后只做报告里点了名、
且严重级允许自动应用的改动 —— 于是「报告说第 5 行重复」与「整理删掉了第 5 行」
是同一句话的两面，不存在两套判定各自跑偏的可能。

**台账的坐标是原文件的坐标。** 输出的行号已经漂了（第 4 行被删之后，
原来的第 5 行变成第 4 行），拿输出的行号去核对等于让用户在文件里找一个不存在的行。
所以 `ChangeLog` 里记的是「输出第 i 行来自原文件第 rows_kept[i] 行」，
以及每一处改动的原坐标与 before/after。

**台账必须完整。** 少记一处改动，用户会以为那个值原样保留着；
多记一处，用户会去核对一个根本不存在的差异。这条由测试里那个不认规则的
对账器守着：拿「台账 + 输入」逐格还原输出，必须完全一致。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook

from . import inspector, rules, values

#: 输出文件名的后缀。与 `excel.write_logistics_result` 的 `_物流信息` 一个路数 ——
#: 用户在自己的文件夹里能一眼看出哪份是整理过的
_SUFFIX = "_已整理"

#: 这些 kind 的动作是「删掉整行」
_ROW_REMOVALS = ("leading_blank_rows", "blank_row", "blank_key", "duplicate_row", "duplicate_key")
#: 这些 kind 的动作是「去掉格子里两端多余的空白」
_CELL_TRIMMINGS = ("header_whitespace", "key_whitespace")


@dataclass(frozen=True)
class CleanOptions:
    #: 主键列。空表示用户还没选 —— 主键相关的规则整体不参与
    key_col: str = ""
    #: 是否应用 `risk` 级规则（主键重复只留第一行）。
    #: 默认关：同一条工单做两遍 = 重复退款，这种事不能替用户决定。
    accept_risk: bool = False


DEFAULT_OPTIONS = CleanOptions()


@dataclass(frozen=True)
class Change:
    """台账里的一处改动。

    `row` / `col` 都是**原文件**的坐标，0 表示这一处不涉及行/列。
    改值的改动带上 `before` / `after`；删行删列没有 after（东西不在了）。
    原因文案跟着 `detail` 一起记，台账导出成 CSV 之后仍然自己说得清楚。
    """

    kind: str
    detail: str
    row: int = 0
    col: int = 0
    before: Any = None
    after: Any = None


@dataclass
class ChangeLog:
    source: str
    output: str
    sheet: str
    #: 输出第 i 行来自原文件的哪一行（严格递增）
    rows_kept: tuple[int, ...]
    #: 输出第 j 列来自原文件的哪一列（严格递增）
    cols_kept: tuple[int, ...]
    changes: tuple[Change, ...] = field(default_factory=tuple)


def clean(path: str, options: CleanOptions = DEFAULT_OPTIONS) -> tuple[str, ChangeLog]:
    """把 `path` 整理成 `原名_已整理.xlsx`，返回 (输出路径, 台账)。

    只应用体检报告里 `fix` 级的改动；`risk` 级要 `options.accept_risk` 才动；
    `info` 级（列名重复、类型认不出来）一个字节都不改 —— 猜错的值会变成客户投诉。
    """
    report = inspector.inspect(path, options.key_col)
    sheet, grid = inspector.read_grid(path)

    live = [
        issue
        for issue in report.issues
        if issue.severity == rules.AUTO_FIX or (options.accept_risk and issue.severity == rules.NEEDS_CONSENT)
    ]

    drop_rows: set[int] = set()
    drop_cols: set[int] = set()
    trim: dict[tuple[int, int], inspector.Issue] = {}
    for issue in live:
        if issue.kind in _ROW_REMOVALS:
            drop_rows.update(issue.rows)
        elif issue.kind == "blank_header":
            drop_cols.update(issue.cols)
        elif issue.kind in _CELL_TRIMMINGS:
            for row in issue.rows:
                for col in issue.cols:
                    trim[(row, col)] = issue

    # 表头行不在任何删除规则里，会自然留下；表头之上的空行由 `leading_blank_rows` 收进 drop_rows
    rows_kept = tuple(r for r in range(1, len(grid) + 1) if r not in drop_rows)
    cols_kept = tuple(c for c in range(1, len(report.columns) + 1) if c not in drop_cols)

    out_grid, cell_changes = _build_grid(grid, rows_kept, cols_kept, trim)
    removal_changes = _removal_changes(live)

    output = output_name(path)
    _write(output, sheet, out_grid)

    log = ChangeLog(
        source=str(path),
        output=output,
        sheet=sheet,
        rows_kept=rows_kept,
        cols_kept=cols_kept,
        changes=tuple(removal_changes + cell_changes),
    )
    return output, log


def _build_grid(
    grid: list[list[Any]],
    rows_kept: tuple[int, ...],
    cols_kept: tuple[int, ...],
    trim: dict[tuple[int, int], inspector.Issue],
) -> tuple[list[list[Any]], list[Change]]:
    """按保留的行列取格子，顺手记下改过值的那些。

    改值的记录**在这一趟里生成**，所以它不可能与真正写进文件的值不一致 ——
    先改文件、再回头补台账的写法，迟早会漏掉一处。
    """
    out: list[list[Any]] = []
    changes: list[Change] = []
    for row_no in rows_kept:
        source_row = grid[row_no - 1]
        line: list[Any] = []
        for col_no in cols_kept:
            before = source_row[col_no - 1] if col_no - 1 < len(source_row) else None
            issue = trim.get((row_no, col_no))
            after = values.normalize_text(before) if issue else before
            if issue is not None and after != before:
                changes.append(
                    Change(kind=issue.kind, detail=issue.detail, row=row_no, col=col_no, before=before, after=after)
                )
            line.append(after)
        out.append(line)
    return out, changes


def _removal_changes(live: list[inspector.Issue]) -> list[Change]:
    """删行删列的记录。按报告的顺序（也就是规则的执行顺序）产出，台账读起来与报告同序。"""
    out: list[Change] = []
    for issue in live:
        if issue.kind in _ROW_REMOVALS:
            out += [Change(kind=issue.kind, detail=issue.detail, row=row) for row in issue.rows]
        elif issue.kind == "blank_header":
            out += [Change(kind=issue.kind, detail=issue.detail, col=col) for col in issue.cols]
    return out


def output_name(path: str) -> str:
    """`订单.xlsx` -> `订单_已整理.xlsx`。

    放在原文件旁边而不是别处：用户是从自己的文件夹里挑的表，结果就该出现在同一个文件夹里，
    否则他得去别的地方找那份文件。
    """
    return re.sub(r"(\.[^.]+)$", rf"{_SUFFIX}\1", path)


def _write(path: str, sheet: str, grid: list[list[Any]]) -> None:
    """只写有值的格子。

    空单元格不写：`ws.cell(..., value=None)` 会把「从没被碰过」变成一个真实的空单元格，
    表格软件里会出现无边界的空白区域，用户选中整列时看到一片假空行。
    """
    wb = Workbook()
    try:
        ws = wb.active
        ws.title = sheet
        for i, row in enumerate(grid, start=1):
            for j, value in enumerate(row, start=1):
                if value is not None:
                    ws.cell(row=i, column=j, value=value)
        wb.save(path)
    finally:
        wb.close()
